"""
Converts the submission Markdown into a Word document.

  python scripts/md-to-docx.py docs/DEXLESS-BNB-SUBMISSION.md export/out.docx

Deliberately hand-rolled rather than pandoc: the document is a small, known set
of constructs (headings, tables, fenced code, blockquotes, lists, links) and a
targeted converter gives control over the things that matter here — a CJK font
that actually resolves in Word, real hyperlinks on the BscScan addresses, and
monospace runs for contract data.
"""

import io
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

BODY_FONT = "Microsoft JhengHei"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x7A, 0x5A, 0x0E)
MUTED = RGBColor(0x55, 0x4E, 0x42)
INK = RGBColor(0x14, 0x10, 0x0A)

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
TOKEN_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def set_cjk(run, font=BODY_FONT):
    """python-docx only sets the Latin face; Word needs eastAsia set too."""
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)


def add_hyperlink(paragraph, url, text, mono=False):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "7A5A0E")
    rPr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    rFonts = OxmlElement("w:rFonts")
    face = MONO_FONT if mono else BODY_FONT
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rFonts.set(qn(attr), face)
    rPr.append(rFonts)

    if mono:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "15")  # half-points
        rPr.append(sz)

    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def write_inline(paragraph, text, size=None, color=None):
    """Render **bold**, `code` and [links](url) inside one paragraph."""
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        m = LINK_RE.fullmatch(part)
        if m:
            label, url = m.group(1), m.group(2)
            label = label.strip("`")
            add_hyperlink(paragraph, url, label, mono=label.startswith("0x"))
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
            set_cjk(r)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            set_cjk(r, MONO_FONT)
            r.font.size = Pt(9)
        else:
            r = paragraph.add_run(part)
            set_cjk(r)
        if size:
            r.font.size = size
        if color:
            r.font.color.rgb = color


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def main(src, dst):
    lines = io.open(src, encoding="utf-8").read().split("\n")
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    # A blank default template has no paragraphs to drop.
    if doc.paragraphs:
        doc.paragraphs[0]._p.getparent().remove(doc.paragraphs[0]._p)

    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.8)
        s.top_margin = s.bottom_margin = Inches(0.75)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ── fenced code ──────────────────────────────────────────
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("\n".join(buf))
            set_cjk(r, MONO_FONT)
            r.font.size = Pt(8.5)
            r.font.color.rgb = INK
            continue

        # ── tables ───────────────────────────────────────────────
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1

            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                write_inline(p, h, size=Pt(8.5))
                for r in p.runs:
                    r.bold = True
                shade(cell, "F0EDE4")
            for row in rows:
                cells = t.add_row().cells
                for j, val in enumerate(row[: len(header)]):
                    cells[j].text = ""
                    write_inline(cells[j].paragraphs[0], val.replace("<br>", " · "), size=Pt(8.5))
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # ── headings ─────────────────────────────────────────────
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(6)
            sizes = {1: 19, 2: 14, 3: 11.5, 4: 10.5}
            write_inline(p, text, size=Pt(sizes.get(level, 10)))
            for r in p.runs:
                r.bold = True
                if level <= 2:
                    r.font.color.rgb = INK
                else:
                    r.font.color.rgb = ACCENT
            i += 1
            continue

        # ── horizontal rule ──────────────────────────────────────
        if stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "C9C2B4")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # ── blockquote ───────────────────────────────────────────
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            for b in buf:
                if not b:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.space_after = Pt(3)
                if b.startswith("###"):
                    b = b.lstrip("#").strip()
                    write_inline(p, b, size=Pt(11))
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = ACCENT
                elif b.startswith("- "):
                    p.style = "List Bullet"
                    p.paragraph_format.left_indent = Inches(0.45)
                    write_inline(p, b[2:], size=Pt(9.5))
                else:
                    write_inline(p, b, size=Pt(9.5))
            continue

        # ── list item ────────────────────────────────────────────
        if re.match(r"^[-*] ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            write_inline(p, stripped[2:])
            i += 1
            continue

        # ── blank ────────────────────────────────────────────────
        if not stripped:
            i += 1
            continue

        # ── paragraph (join wrapped lines) ───────────────────────
        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "```", "---"))
                    or re.match(r"^[-*] ", nxt)):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        write_inline(p, " ".join(buf))

    doc.save(dst)
    print(f"wrote {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
